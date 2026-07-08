from __future__ import annotations

from pathlib import Path


def validate_input_headers(file_path: Path, required_headers: list[str]) -> tuple[bool, str]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError:
        return False, "Dependência ausente: instale openpyxl antes da validação de Excel."

    if not file_path.exists():
        return False, f"Arquivo não encontrado: {file_path}"

    workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    sheet = workbook.active
    first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = [str(cell).strip().lower() for cell in first_row if cell is not None]
    workbook.close()

    missing = [header for header in required_headers if header.lower() not in headers]
    if missing:
        return False, f"Colunas obrigatórias ausentes: {', '.join(missing)}"
    return True, "OK"
